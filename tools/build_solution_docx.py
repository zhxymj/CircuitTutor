# -*- coding: utf-8 -*-
"""从作品说明书 Markdown 生成正式企业命题 DOCX。"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from core.version import VERSION as PRODUCT_VERSION

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
except ModuleNotFoundError as exc:
    if exc.name != "docx":
        raise
    # Version and release metadata can be inspected without the optional
    # document-generation dependency. Building still reports the missing tool.
    Document = None


TITLE = "智绘电路 面向通信设备低压 MCU 与 I2C 传感接口子板的 AI 候选约束与原理图工程验证系统"
SUBTITLE = "中兴通讯 AI辅助硬件设计 企业命题 命题1550"


def set_font(run, name="等线", size=10.5, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_table_row_split(row):
    """Keep a logical table row on one page to avoid orphaned cell fragments."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def remove_paragraph_border(paragraph_or_style):
    """Remove built-in Title style borders that render as an unwanted rule."""
    element = paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    return run


def add_inline(paragraph, text, default_size=10.5):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_font(paragraph.add_run(text[pos:match.start()]), size=default_size)
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), size=default_size, bold=True)
        else:
            set_font(paragraph.add_run(token[1:-1]), name="Consolas", size=max(8.5, default_size - 1))
        pos = match.end()
    if pos < len(text):
        set_font(paragraph.add_run(text[pos:]), size=default_size)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "等线"
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = "黑体"
    title._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(16)
    remove_paragraph_border(title)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "等线"
    subtitle._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "等线")
    subtitle.font.size = Pt(13)
    subtitle.font.italic = False
    subtitle.font.color.rgb = RGBColor(0, 0, 0)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name, size, font, before, after in (
        ("Heading 1", 16, "黑体", 18, 8),
        ("Heading 2", 13.5, "黑体", 13, 6),
        ("Heading 3", 11.5, "黑体", 10, 4),
    ):
        style = styles[name]
        style.font.name = font
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    if "Source Code" not in styles:
        code = styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Source Code"]
    code.font.name = "Consolas"
    code._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "等线")
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor(0, 0, 0)
    code.paragraph_format.left_indent = Cm(0.45)
    code.paragraph_format.right_indent = Cm(0.25)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.05


def configure_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.45)
    section.bottom_margin = Cm(2.35)
    section.left_margin = Cm(2.65)
    section.right_margin = Cm(2.45)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def add_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(paragraph.add_run("中兴通讯 AI辅助硬件设计 企业命题解决方案    "), size=8.5, color="666666")
    run = add_field(paragraph, "PAGE", "1")
    set_font(run, size=8.5, color="666666")


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    remove_paragraph_border(p)
    p.paragraph_format.space_before = Pt(32)
    p.add_run(TITLE)
    p = doc.add_paragraph(style="Subtitle")
    p.add_run(SUBTITLE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    set_font(p.add_run("技术解决书"), name="黑体", size=18, bold=True)
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("A 主线  AI 辅助原理图设计合规检查与工程验证"), size=11.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("B 方向  PCB 智能布线接口预研"), size=11.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(54)
    set_font(p.add_run("2026 年 9 月 19 日"), size=11)
    doc.add_page_break()


def add_toc(doc):
    heading = doc.add_paragraph("目录", style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "目录将在 Microsoft Word 中更新")
    doc.add_page_break()


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        return [rows[0]] + rows[2:], i
    return None, start


def add_table(doc, rows):
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    repeat_table_header(table.rows[0])
    page_width = 15.9
    maxima = []
    for col in range(cols):
        maxima.append(max(4, max(len(row[col]) if col < len(row) else 0 for row in rows)))
    total = sum(maxima)
    widths = [max(1.5, page_width * value / total) for value in maxima]
    scale = page_width / sum(widths)
    widths = [value * scale for value in widths]
    font_size = 8.2 if cols >= 6 else 8.8 if cols >= 5 else 9.2
    for r_idx, row in enumerate(rows):
        prevent_table_row_split(table.rows[r_idx])
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.width = Cm(widths[c_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.05
            value = row[c_idx] if c_idx < len(row) else ""
            add_inline(paragraph, value, default_size=font_size)
            if r_idx == 0:
                set_cell_shading(cell, "1F4E78")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if r_idx % 2 == 0:
                    set_cell_shading(cell, "F3F6F9")
                if len(value) <= 14 and c_idx > 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_body(doc, markdown):
    lines = markdown.splitlines()
    if lines and lines[0].startswith("# "):
        lines = lines[1:]
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph(style="Source Code")
                p.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines):
            table_rows, next_i = parse_table(lines, i)
            if table_rows:
                add_table(doc, table_rows)
                i = next_i
                continue
        heading_match = re.match(r"^(#{2,4})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1)) - 1
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            add_inline(p, heading_match.group(2), default_size={1: 16, 2: 13.5, 3: 11.5}[min(level, 3)])
            i += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph(style="Quote")
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.right_indent = Cm(0.35)
            add_inline(p, stripped[2:])
            i += 1
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
            i += 1
            continue
        if numbered:
            # Preserve the source list number explicitly. Word's built-in List
            # Number style otherwise continues numbering across independent lists.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            set_font(p.add_run(f"{numbered.group(1)}.  "), bold=True)
            add_inline(p, numbered.group(2))
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1


def build(source, output):
    if Document is None:
        raise RuntimeError("生成 DOCX 需要安装 python-docx")
    text = Path(source).read_text(encoding="utf-8")
    doc = Document()
    configure_section(doc.sections[0])
    configure_styles(doc)
    add_footer(doc.sections[0])
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "中兴通讯 AI辅助硬件设计 企业命题技术解决书"
    doc.core_properties.author = "参赛团队"
    doc.core_properties.version = PRODUCT_VERSION
    doc.settings.element.append(OxmlElement("w:updateFields"))
    doc.settings.element[-1].set(qn("w:val"), "true")
    add_cover(doc)
    add_toc(doc)
    add_body(doc, text)
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source")
    parser.add_argument("output")
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()
